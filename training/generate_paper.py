import os
import subprocess
import sys

# Auto-install python-docx if missing
try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

def generate_editable_doc():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading("Adaptive Residual Learning for Mobile Image Denoising: A Convolutional Approach", 0)
    title.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    
    # Abstract
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "Smartphone camera sensors are inherently limited by their physical size, making them highly susceptible to "
        "high-frequency photon and read noise, particularly in low-light environments. This paper presents a Phase 3 "
        "image denoising pipeline utilizing a 17-layer Deep Convolutional Neural Network (DnCNN). By employing additive "
        "residual learning, the model isolates and extracts structural noise maps rather than predicting clean images "
        "directly. Evaluated against the Smartphone Image Denoising Dataset (SIDD), the architecture demonstrated a 10.5% "
        "noise reduction on a 160-pair baseline, which scaled to a 28.3% reduction when trained on a broader 320-pair "
        "medium dataset. The final computational graph is exported via ONNX to facilitate LiteRT mobile deployment."
    )
    
    # 1. Introduction
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Traditional Image Signal Processors (ISPs) rely on static filters that often compromise fine texture details "
        "when aggressively smoothing sensor grain. Neural network-based denoising provides a data-driven alternative, "
        "identifying complex spatial noise patterns across varying illumination levels. This research outlines the "
        "development, training, and deployment preparation of a convolutional denoising module designed to integrate "
        "into a multi-stage mobile computational photography pipeline."
    )
    
    # 2. Architecture and Methodology
    doc.add_heading("2. Architecture and Methodology", level=1)
    doc.add_paragraph(
        "The core network utilizes a 17-layer DnCNN architecture stripped of dense (fully connected) layers, ensuring "
        "the system remains completely resolution-agnostic during both local execution and eventual mobile inference."
    )
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Residual Learning Formulation: ").bold = True
    p.add_run("The network maps the transformation y = x + v, where y is the observed noisy input, x is the latent "
              "clean image, and v is the additive noise. The network is trained to predict v, which is subsequently "
              "subtracted from y.")
    p2 = doc.add_paragraph(style='List Bullet')
    p2.add_run("Hardware Constraints: ").bold = True
    p2.add_run("To operate within a strict 6GB VRAM hardware limitation (NVIDIA RTX 3050), memory buffers were "
               "optimized utilizing dynamic spatial cropping, processing 128x128 image patches at a batch size of 16.")
    
    # 3. Dataset and Domain Processing
    doc.add_heading("3. Dataset and Domain Processing", level=1)
    doc.add_paragraph(
        "The model was trained utilizing the Smartphone Image Denoising Dataset (SIDD). The medium tier of this dataset "
        "contains 320 image pairs representing 160 distinct scene instances captured across five smartphone models."
    )
    p3 = doc.add_paragraph(style='List Bullet')
    p3.add_run("Tensor Normalization: ").bold = True
    p3.add_run("Standard 8-bit gamma-corrected sRGB integer files (.PNG) were ingested and dynamically scaled into "
               "precise float32 tensors representing a [0.0, 1.0] numerical scale.")
    p4 = doc.add_paragraph(style='List Bullet')
    p4.add_run("Linear Float Integration: ").bold = True
    p4.add_run("To accommodate target mobile FusedFrame data contracts, an auxiliary pipeline was engineered using the "
               "rawpy library. This allows the system to bypass conventional ISP gamma compression, extracting uncompressed "
               "16-bit linear sensor data directly from RAW (.dng) files.")
    
    # 4. Experimental Results
    doc.add_heading("4. Experimental Results and Analytics", level=1)
    doc.add_paragraph(
        "Because custom camera test captures lack perfect ground-truth references, the model’s efficacy was evaluated "
        "using Immerkaer’s Fast Noise Variance formula, utilizing PyTorch-native Laplacian convolutional filters to "
        "mathematically estimate the standard deviation (σ) of high-frequency pixel grain."
    )
    
    # Add Table for Results
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Training Dataset'
    hdr_cells[1].text = 'Original Noise (σ)'
    hdr_cells[2].text = 'Denoised Noise (σ)'
    hdr_cells[3].text = 'Reduction (%)'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'SIDD Small (160 Pairs)'
    row_cells[1].text = '3.86'
    row_cells[2].text = '3.45'
    row_cells[3].text = '10.5%'
    
    row_cells2 = table.add_row().cells
    row_cells2[0].text = 'SIDD Medium (320 Pairs)'
    row_cells2[1].text = '3.86'
    row_cells2[2].text = '2.77'
    row_cells2[3].text = '28.3%'
    
    doc.add_paragraph("") # Spacing
    
    # 5. Mobile Deployment
    doc.add_heading("5. Mobile Deployment and Handoff", level=1)
    doc.add_paragraph(
        "To bridge the PyTorch research environment with the downstream mobile LiteRT (TFLite) deployment requirements, "
        "the dynamic weights were frozen and exported into a strict .onnx computational graph. An automated cryptographic "
        "manifest generator was built to output a deployment payload comprising the model file, its SHA-256 hash, Git "
        "commit tracking, and explicit input domain declarations."
    )
    
    # 6. Conclusion
    doc.add_heading("6. Conclusion and Future Work", level=1)
    doc.add_paragraph(
        "The 17-layer DnCNN pipeline successfully demonstrated scalable noise extraction capabilities in the sRGB domain. "
        "Future work will focus on addressing domain shift by loading the pre-trained weights and fine-tuning the network "
        "explicitly on linear fusion-residual data pairs provided by upstream alignment modules, optimizing the model for "
        "real-world raw sensor data on device."
    )
    
    # Save the document
    filename = "Role4_Denoising_Research_Paper.docx"
    doc.save(filename)
    print(f"✅ Document generated successfully: {os.path.abspath(filename)}")

if __name__ == "__main__":
    generate_editable_doc()